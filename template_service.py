
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

def build_input_templates(output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    inscritos = output_dir / 'plantilla_inscritos.xlsx'
    instrumento = output_dir / 'plantilla_instrumento.xlsx'

    wb = Workbook()
    ws = wb.active
    ws.title = 'inscritos'
    headers = ['rut','nombre_completo','email','curso','establecimiento','fecha_inscripcion','consentimiento_correo','telefono','interes_carreras']
    ws.append(headers)
    ws.append(['12345678-5','Nombre Ejemplo','correo@ejemplo.cl','4 Medio','Colegio Ejemplo','2026-01-01','SI','',''])
    table = Table(displayName='inscritos_tbl', ref=f'A1:I2')
    style = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    ws.add_table(table)
    wb.save(inscritos)

    wb = Workbook()
    ws = wb.active
    ws.title = 'instrumento'
    headers = ['pregunta','respuesta_correcta','puntaje','item_id','eje','habilidad','unidad_tematica','nivel_dificultad','usa_en_puntaje_paes']
    ws.append(headers)
    ws.append([1,'A',1,'I1','Números','Resolver','Conjunto de los números reales','Media','SI'])
    table = Table(displayName='instrumento_tbl', ref=f'A1:I2')
    table.tableStyleInfo = style
    ws.add_table(table)
    wb.save(instrumento)
    return inscritos, instrumento
